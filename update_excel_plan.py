import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb_path = 'd:/cash_machine/Chathura_Printers_Business_Plan.xlsx'
wb = openpyxl.load_workbook(wb_path)

# If sheet already exists, remove it so we can create it fresh
sheet_name = "What's Left at Month End"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws = wb.create_sheet(title=sheet_name, index=1)

# Color styles
FONT_NAME = 'Arial'
title_font = Font(name=FONT_NAME, size=14, bold=True, color='1E293B')
subtitle_font = Font(name=FONT_NAME, size=9.5, italic=True, color='64748B')
section_header_font = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
header_font = Font(name=FONT_NAME, size=10, bold=True, color='0F172A')
bold_font = Font(name=FONT_NAME, size=10, bold=True, color='0F172A')
regular_font = Font(name=FONT_NAME, size=10, color='1E293B')
italic_font = Font(name=FONT_NAME, size=9, italic=True, color='64748B')
input_font = Font(name=FONT_NAME, size=10, bold=True, color='0000FF') # Blue for user inputs
kpi_title_font = Font(name=FONT_NAME, size=8.5, bold=True, color='475569')
kpi_num_font = Font(name=FONT_NAME, size=12, bold=True, color='0F172A')
kpi_num_green = Font(name=FONT_NAME, size=12, bold=True, color='047857')

section_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # Slate 800
table_header_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid') # Slate 100
input_fill = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid') # Light blue fill
highlight_fill = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid') # Emerald 50
warning_fill = PatternFill(start_color='FFFBEB', end_color='FFFBEB', fill_type='solid') # Amber 50
kpi_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

thin_border_side = Side(style='thin', color='CBD5E1')
double_bottom_side = Side(style='double', color='0F172A')
thick_bottom_side = Side(style='medium', color='0F172A')

border_all_thin = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_total = Border(top=thin_border_side, bottom=double_bottom_side)
border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)

# Number formats
FMT_CURRENCY = '#,##0;(#,##0);"-"'
FMT_PERCENT = '0.0%'
FMT_DECIMAL = '0.00'

# Title
ws['A1'] = "Chathura Printers — What’s Left at Month End & Bank Loan Calculator"
ws['A1'].font = title_font
ws['A2'] = "Legend: Blue bold numbers = editable inputs you can change; Black = formula calculated; Green = verified/confirmed status."
ws['A2'].font = subtitle_font

# -------------------------------------------------------------
# EXECUTIVE KPI SUMMARY BLOCK (Rows 4-7)
# -------------------------------------------------------------
ws['A4'] = "EXECUTIVE SUMMARY (Sample Month at Rs. 460,000 Revenue)"
ws['A4'].font = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
ws['A4'].fill = section_fill
ws.merge_cells('A4:F4')

kpis = [
    ("Monthly Revenue", "=C36", FMT_CURRENCY, kpi_num_font, "A", "B"),
    ("Fixed Expenses", "=B29", FMT_CURRENCY, kpi_num_font, "C", "C"),
    ("Monthly Loan Installment", "=B20", FMT_CURRENCY, kpi_num_font, "D", "D"),
    ("Left to Reinvest (Surplus)", "=C42", FMT_CURRENCY, kpi_num_green, "E", "E"),
    ("Break-Even Monthly Rev.", "=C44", FMT_CURRENCY, kpi_num_font, "F", "F"),
]

for label, val_formula, num_fmt, f_font, col_start, col_end in kpis:
    c_label = ws[f'{col_start}5']
    c_label.value = label
    c_label.font = kpi_title_font
    c_label.alignment = Alignment(horizontal='center', vertical='center')
    c_label.fill = kpi_fill
    
    c_val = ws[f'{col_start}6']
    c_val.value = val_formula
    c_val.font = f_font
    c_val.number_format = num_fmt
    c_val.alignment = Alignment(horizontal='center', vertical='center')
    c_val.fill = kpi_fill
    
    if col_start != col_end:
        ws.merge_cells(f'{col_start}5:{col_end}5')
        ws.merge_cells(f'{col_start}6:{col_end}6')

for r in [5, 6]:
    for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws[f'{col_l}{r}'].border = border_all_thin

# -------------------------------------------------------------
# SECTION 1: FIXED PERSONAL LOAN CALCULATOR & BANK VERIFICATION (Rows 8-22)
# -------------------------------------------------------------
ws['A8'] = "1. FIXED PERSONAL LOAN CALCULATOR & BANK CONFIRMATION"
ws['A8'].font = section_header_font
ws['A8'].fill = section_fill
ws.merge_cells('A8:F8')

ws['A9'] = "Loan Parameter (Inputs)"
ws['B9'] = "Your Value"
ws['C9'] = "Notes / Assumptions"
ws['D9'] = "Calculated Metric"
ws['E9'] = "Formula Result (Rs.)"
ws['F9'] = "Calculation Details"

for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
    cell = ws[f'{col_l}9']
    cell.font = bold_font
    cell.fill = table_header_fill
    cell.border = border_header
    cell.alignment = Alignment(horizontal='left' if col_l in ['A', 'C', 'D', 'F'] else 'right')

loan_rows = [
    # (A, B, C, D, E, F, B_fmt, E_fmt, is_input)
    ("Loan amount / principal (Rs.)", 2800000, "Rs. 28 lakh personal loan", "Monthly installment (EMI)", "=PMT(B11/12,B13,-B10)", "Standard reducing-balance monthly payment", FMT_CURRENCY, FMT_CURRENCY, True),
    ("Annual fixed interest rate", 0.155, "15.5% annual fixed rate (bank quote)", "Initial month interest (M1)", "=B10*(B11/12)", "= Principal × Rate / 12", FMT_PERCENT, FMT_CURRENCY, True),
    ("Loan term (years)", 5, "5 years tenure", "Average monthly interest", "=(E14-B10)/B13", "Total interest divided over all months", "#,##0", FMT_CURRENCY, True),
    ("Total repayment tenure (months)", "=B12*12", "= Years × 12 months", "Total repayment over term", "=E10*B13", "Total amount paid (Principal + Interest)", "#,##0", FMT_CURRENCY, False),
    ("Working days per month", 26, "Used for daily break-even", "Total interest payable", "=E13-B10", "Total finance charge over full tenure", "#,##0", FMT_CURRENCY, True),
    ("Flat Rate: Total flat interest", "=B10*B11*B12", "= Principal × Rate × Years", "Flat monthly installment", "=(B10+A15_interest)/B13", "Alternative: if bank charges flat rate", FMT_CURRENCY, FMT_CURRENCY, False),
    ("Flat Rate: Total repayment", "=B10+B15", "Principal + Flat Interest", "Flat monthly interest", "=B15/B13", "Flat monthly interest portion", FMT_CURRENCY, FMT_CURRENCY, False),
]

row_idx = 10
# Row 10
ws['A10'] = "Loan principal amount (Rs.)"
ws['B10'] = 2800000
ws['B10'].font = input_font
ws['B10'].fill = input_fill
ws['B10'].number_format = FMT_CURRENCY
ws['C10'] = "Rs. 28 lakh principal (change as needed)"
ws['D10'] = "Monthly installment (Reducing EMI)"
ws['E10'] = "=PMT(B11/12,B13,-B10)"
ws['E10'].font = bold_font
ws['E10'].number_format = FMT_CURRENCY
ws['F10'] = "Standard bank reducing balance payment"

# Row 11
ws['A11'] = "Annual fixed interest rate"
ws['B11'] = 0.155
ws['B11'].font = input_font
ws['B11'].fill = input_fill
ws['B11'].number_format = FMT_PERCENT
ws['C11'] = "15.5% fixed interest rate per annum"
ws['D11'] = "Month 1 interest amount"
ws['E11'] = "=B10*(B11/12)"
ws['E11'].font = regular_font
ws['E11'].number_format = FMT_CURRENCY
ws['F11'] = "First month interest portion (= P × r / 12)"

# Row 12
ws['A12'] = "Loan term (years)"
ws['B12'] = 5
ws['B12'].font = input_font
ws['B12'].fill = input_fill
ws['B12'].number_format = "0"
ws['C12'] = "5 years period"
ws['D12'] = "Average monthly interest"
ws['E12'] = "=(E14-B10)/B13"
ws['E12'].font = regular_font
ws['E12'].number_format = FMT_CURRENCY
ws['F12'] = "Average interest per month over tenure"

# Row 13
ws['A13'] = "Loan term (months)"
ws['B13'] = "=B12*12"
ws['B13'].font = bold_font
ws['B13'].number_format = "0"
ws['C13'] = "= Years × 12 months"
ws['D13'] = "Total interest paid over term"
ws['E13'] = "=E14-B10"
ws['E13'].font = regular_font
ws['E13'].number_format = FMT_CURRENCY
ws['F13'] = "Total interest cost across entire 5 years"

# Row 14
ws['A14'] = "Working days per month"
ws['B14'] = 26
ws['B14'].font = input_font
ws['B14'].fill = input_fill
ws['B14'].number_format = "0"
ws['C14'] = "26 working days (Sri Lanka printing shop)"
ws['D14'] = "Total loan repayment over term"
ws['E14'] = "=E10*B13"
ws['E14'].font = bold_font
ws['E14'].number_format = FMT_CURRENCY
ws['F14'] = "Total money paid to bank (Principal + Interest)"

# Row 15: Flat Rate comparison
ws['A15'] = "Flat Rate Alternative: Total Interest"
ws['B15'] = "=B10*B11*B12"
ws['B15'].font = regular_font
ws['B15'].number_format = FMT_CURRENCY
ws['C15'] = "= Principal × Rate × Years (Flat method)"
ws['D15'] = "Flat Rate: Monthly Installment"
ws['E15'] = "=(B10+B15)/B13"
ws['E15'].font = regular_font
ws['E15'].number_format = FMT_CURRENCY
ws['F15'] = "Installment if bank uses flat rate instead of EMI"

for r in range(10, 16):
    for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
        cell = ws[f'{col_l}{r}']
        cell.border = border_all_thin
        if col_l in ['A', 'C', 'D', 'F'] and not cell.font:
            cell.font = regular_font
        if col_l in ['B', 'E']:
            cell.alignment = Alignment(horizontal='right')

# BANK CONFIRMATION / VERIFICATION BLOCK (Rows 17-21)
ws['A17'] = "BANK QUOTE CONFIRMATION & VERIFICATION"
ws['A17'].font = Font(name=FONT_NAME, size=10, bold=True, color='047857')
ws['A17'].fill = highlight_fill
ws.merge_cells('A17:F17')
ws['A17'].alignment = Alignment(horizontal='left')
ws['A17'].border = border_all_thin

# Row 18
ws['A18'] = "Bank Quoted Monthly Installment (Rs.)"
ws['B18'] = 67350
ws['B18'].font = input_font
ws['B18'].fill = input_fill
ws['B18'].number_format = FMT_CURRENCY
ws['B18'].alignment = Alignment(horizontal='right')
ws['C18'] = "Enter the exact monthly installment quoted by your bank manager"

ws['D18'] = "Installment Variance vs Calculated EMI"
ws['E18'] = "=B18-E10"
ws['E18'].font = bold_font
ws['E18'].number_format = 'Rs. #,##0.00;[Red](Rs. #,##0.00);Rs. 0.00'
ws['E18'].alignment = Alignment(horizontal='right')
ws['F18'] = "= Bank Quote − Calculated Reducing EMI (Goal is ~Rs. 0)"

# Row 19
ws['A19'] = "Bank Quoted Monthly Interest Amount (Rs.)"
ws['B19'] = 36167
ws['B19'].font = input_font
ws['B19'].fill = input_fill
ws['B19'].number_format = FMT_CURRENCY
ws['B19'].alignment = Alignment(horizontal='right')
ws['C19'] = "Optional: Month 1 interest quoted by bank (matches ~Rs. 36,167)"

ws['D19'] = "Interest Variance vs Month 1 Interest"
ws['E19'] = "=IF(B19=0,0,B19-E11)"
ws['E19'].font = regular_font
ws['E19'].number_format = 'Rs. #,##0.00;[Red](Rs. #,##0.00);Rs. 0.00'
ws['E19'].alignment = Alignment(horizontal='right')
ws['F19'] = "= Bank Quote − Month 1 Interest (P × r / 12)"

# Row 20
ws['A20'] = "Active Loan Installment Used in Cash Flow"
ws['B20'] = "=IF(B18>0,B18,E10)"
ws['B20'].font = Font(name=FONT_NAME, size=11, bold=True, color='047857')
ws['B20'].fill = highlight_fill
ws['B20'].number_format = FMT_CURRENCY
ws['B20'].alignment = Alignment(horizontal='right')
ws['C20'] = "Uses confirmed bank quote (or calculated EMI if blank)"

ws['D20'] = "Confirmation Status"
ws['E20'] = '=IF(B18=0,"Enter Bank Quote",IF(ABS(E18)<=10,"CONFIRMED: MATCHES BANK QUOTE",IF(ABS(B18-E15)<=10,"MATCHES FLAT RATE QUOTE",IF(E18>0,"BANK QUOTE IS HIGHER","BANK QUOTE IS LOWER"))))'
ws['E20'].font = Font(name=FONT_NAME, size=10, bold=True, color='047857')
ws['E20'].fill = highlight_fill
ws['E20'].alignment = Alignment(horizontal='center')
ws.merge_cells('E20:F20')

for r in range(18, 21):
    for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws[f'{col_l}{r}'].border = border_all_thin

# -------------------------------------------------------------
# SECTION 2: FIXED MONTHLY EXPENSES (Rows 22-30)
# -------------------------------------------------------------
ws['A22'] = "2. FIXED MONTHLY EXPENSES (Updated: No Rent, No Licences/Insurance)"
ws['A22'].font = section_header_font
ws['A22'].fill = section_fill
ws.merge_cells('A22:F22')

ws['A23'] = "Expense Item"
ws['B23'] = "Amount (Rs.)"
ws['C23'] = "Details / Breakdown"
ws['D23'] = "Category / Notes"
ws.merge_cells('D23:F23')

for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
    cell = ws[f'{col_l}23']
    cell.font = bold_font
    cell.fill = table_header_fill
    cell.border = border_header
    cell.alignment = Alignment(horizontal='left' if col_l != 'B' else 'right')

fixed_expenses = [
    ("Manager salary (Rs. 50,000 basic + EPF/ETF Rs. 7,500)", 57500, "12% EPF (Rs. 6,000) + 3% ETF (Rs. 1,500) on Rs. 50k", "Payroll / Staff"),
    ("Electricity", 18000, "Heavy production machinery power consumption", "Utilities"),
    ("Machine repairs and servicing", 8000, "Preventive maintenance, drum & blade replacement", "Maintenance"),
    ("Marketing", 5000, "Flyers, local social media promos, customer canvassing", "Sales / Promo"),
    ("Phone and internet", 4000, "High-speed broadband for customer file uploads + phone line", "Communications"),
]

for idx, (exp_name, exp_amt, exp_desc, exp_cat) in enumerate(fixed_expenses, start=24):
    ws[f'A{idx}'] = exp_name
    ws[f'A{idx}'].font = regular_font
    ws[f'B{idx}'] = exp_amt
    ws[f'B{idx}'].font = input_font
    ws[f'B{idx}'].fill = input_fill
    ws[f'B{idx}'].number_format = FMT_CURRENCY
    ws[f'B{idx}'].alignment = Alignment(horizontal='right')
    ws[f'C{idx}'] = exp_desc
    ws[f'C{idx}'].font = italic_font
    ws[f'D{idx}'] = exp_cat
    ws[f'D{idx}'].font = regular_font
    ws.merge_cells(f'D{idx}:F{idx}')
    for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws[f'{col_l}{idx}'].border = border_all_thin

# Total Fixed Expenses Row (Row 29)
ws['A29'] = "Total Fixed Monthly Expenses"
ws['A29'].font = bold_font
ws['B29'] = "=SUM(B24:B28)"
ws['B29'].font = bold_font
ws['B29'].number_format = FMT_CURRENCY
ws['B29'].alignment = Alignment(horizontal='right')
ws['C29'] = "Strict fixed overhead (Rent & Licences set to Rs. 0 as instructed)"
ws['C29'].font = bold_font
ws['D29'] = "Operational Overhead Base"
ws['D29'].font = bold_font
ws.merge_cells('D29:F29')

for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
    ws[f'{col_l}29'].border = border_total

# -------------------------------------------------------------
# SECTION 3: WHAT'S LEFT AT MONTH END (Rows 32-47)
# -------------------------------------------------------------
ws['A32'] = "3. WHAT’S LEFT AT MONTH END (Revenue Scenarios & Custom Planner)"
ws['A32'].font = section_header_font
ws['A32'].fill = section_fill
ws.merge_cells('A32:F32')

ws['A34'] = "Financial Metric / Cash Flow Step"
ws['B34'] = "Slow Month"
ws['C34'] = "Sample Month"
ws['D34'] = "Busy Month"
ws['E34'] = "Custom Target"
ws['F34'] = "Formula Logic / Strategic Meaning"

for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
    cell = ws[f'{col_l}34']
    cell.font = bold_font
    cell.fill = table_header_fill
    cell.border = border_header
    cell.alignment = Alignment(horizontal='left' if col_l in ['A', 'F'] else 'right')

# Header descriptions row (Row 35)
ws['A35'] = "Sales Volume Level"
ws['B35'] = "Rs. 300,000 / mo"
ws['C35'] = "Rs. 460,000 / mo"
ws['D35'] = "Rs. 600,000 / mo"
ws['E35'] = "Editable Goal"
ws['F35'] = "Scenarios comparison for owner decision making"
for col_l in ['A', 'B', 'C', 'D', 'E', 'F']:
    cell = ws[f'{col_l}35']
    cell.font = italic_font
    cell.border = border_all_thin
    cell.alignment = Alignment(horizontal='left' if col_l in ['A', 'F'] else 'right')

scenario_rows = [
    # (Row_idx, Label, B_val, C_val, D_val, E_val, F_desc, num_fmt, is_bold, is_input)
    (36, "Monthly revenue", 300000, 460000, 600000, 500000, "Gross sales generated in the calendar month", FMT_CURRENCY, True, True),
    (37, "Gross profit margin (%)", "='Monthly Plan'!$F$8", "='Monthly Plan'!$F$8", "='Monthly Plan'!$F$8", "='Monthly Plan'!$F$8", "Overall direct product margin (40.0% from print jobs)", FMT_PERCENT, False, False),
    (38, "Gross profit (Rs.)", "=B36*B37", "=C36*C37", "=D36*D37", "=E36*E37", "= Monthly Revenue × Gross Margin (40%)", FMT_CURRENCY, True, False),
    (39, "Less: Fixed expenses (Rs.)", "=-$B$29", "=-$B$29", "=-$B$29", "=-$B$29", "Manager salary + Electricity + Repairs + Promo + Phone", FMT_CURRENCY, False, False),
    (40, "Operating profit (Rs.)", "=B38+B39", "=C38+C39", "=D38+D39", "=E38+E39", "Operating profit before debt service (= Gross Profit − Overhead)", FMT_CURRENCY, True, False),
    (41, "Less: Loan installment (Rs.)", "=-$B$20", "=-$B$20", "=-$B$20", "=-$B$20", "Fixed monthly bank loan installment (confirmed with bank)", FMT_CURRENCY, False, False),
    (42, "Left to reinvest in business (Rs.)", "=B40+B41", "=C40+C41", "=D40+D41", "=E40+E41", "NET SURPLUS (Positive = reinvest/savings; Negative = top-up)", FMT_CURRENCY, True, False),
    (43, "Cash flow status for owner", '=IF(B42>=0,"SURPLUS: Covers loan + manager","DEFICIT: Top-up from personal salary")', '=IF(C42>=0,"SURPLUS: Covers loan + manager","DEFICIT: Top-up from personal salary")', '=IF(D42>=0,"SURPLUS: Covers loan + manager","DEFICIT: Top-up from personal salary")', '=IF(E42>=0,"SURPLUS: Covers loan + manager","DEFICIT: Top-up from personal salary")', "Owner cash security assessment", "@", True, False),
    (44, "Break-even monthly revenue (Rs.)", "=($B$29+$B$20)/B37", "=($B$29+$B$20)/C37", "=($B$29+$B$20)/D37", "=($B$29+$B$20)/E37", "= (Fixed Expenses + Loan Installment) / Gross Margin", FMT_CURRENCY, True, False),
    (45, "Break-even revenue per working day", "=B44/$B$14", "=C44/$B$14", "=D44/$B$14", "=E44/$B$14", "= Break-even monthly revenue / 26 working days", FMT_CURRENCY, True, False),
]

for r_idx, label, b_val, c_val, d_val, e_val, f_desc, n_fmt, bold_flag, input_flag in scenario_rows:
    ws[f'A{r_idx}'] = label
    ws[f'A{r_idx}'].font = bold_font if bold_flag else regular_font
    
    for c_letter, val in [('B', b_val), ('C', c_val), ('D', d_val), ('E', e_val)]:
        cell = ws[f'{c_letter}{r_idx}']
        cell.value = val
        cell.number_format = n_fmt
        cell.alignment = Alignment(horizontal='right')
        if r_idx == 36:
            cell.font = input_font
            cell.fill = input_fill
        elif r_idx == 42:
            cell.font = Font(name=FONT_NAME, size=11, bold=True, color='047857' if c_letter in ['C', 'D', 'E'] else 'B91C1C')
            cell.fill = highlight_fill if c_letter in ['C', 'D', 'E'] else warning_fill
            cell.border = border_total
        elif r_idx == 43:
            cell.font = Font(name=FONT_NAME, size=8.5, bold=True, color='047857' if c_letter in ['C', 'D', 'E'] else 'B91C1C')
            cell.alignment = Alignment(horizontal='center')
        elif bold_flag:
            cell.font = bold_font
        else:
            cell.font = regular_font

    ws[f'F{r_idx}'] = f_desc
    ws[f'F{r_idx}'].font = italic_font
    
    for c_letter in ['A', 'B', 'C', 'D', 'E', 'F']:
        if r_idx != 42:
            ws[f'{c_letter}{r_idx}'].border = border_all_thin

# STRATEGIC ADVISORY NOTES (Rows 47-51)
ws['A47'] = "OWNER SUMMARY & RISK MANAGEMENT GUIDELINES:"
ws['A47'].font = bold_font
ws.merge_cells('A47:F47')

advisory = [
    "1. Break-even threshold is ~Rs. 400,000 / month (~Rs. 15,385 / day across 26 working days). Any sales above this directly build business cash reserve.",
    "2. In the Sample Month (Rs. 460,000 sales): The business fully covers the manager salary, EPF/ETF, utility bills, and loan installment, leaving +Rs. 24,150 surplus.",
    "3. In a Busy Month (Rs. 600,000 sales): Monthly reinvestment surplus jumps to +Rs. 80,150 to accelerate equipment payoff or store improvements.",
    "4. In a Slow Month (Rs. 300,000 sales): Cash shortfall is -Rs. 39,850. Owner must top up from personal salary or keep a 2-3 month installment cushion (~Rs. 200,000)."
]

for idx, adv in enumerate(advisory, start=48):
    ws[f'A{idx}'] = adv
    ws[f'A{idx}'].font = italic_font
    ws.merge_cells(f'A{idx}:F{idx}')

# Set Column Widths
ws.column_dimensions['A'].width = 46.0
ws.column_dimensions['B'].width = 22.0
ws.column_dimensions['C'].width = 24.0
ws.column_dimensions['D'].width = 24.0
ws.column_dimensions['E'].width = 26.0
ws.column_dimensions['F'].width = 54.0

# -------------------------------------------------------------
# Cross-synchronize existing sheets: 'Inputs', 'Monthly Plan', 'Scenarios'
# -------------------------------------------------------------
# Update Inputs sheet to reference the confirmed loan installment if desired
ws_inputs = wb['Inputs']
ws_inputs['A8'] = "Monthly installment (EMI)"
ws_inputs['C8'] = "Standard reducing-balance loan; check 'What''s Left at Month End' for bank verification"

# Update Scenarios sheet to link to 'What's Left at Month End' active installment and fixed expenses
ws_scenarios = wb['Scenarios']
ws_scenarios['A10'] = "Loan installment"
ws_scenarios['B10'] = "=-'What''s Left at Month End'!$B$20"
ws_scenarios['C10'] = "=-'What''s Left at Month End'!$B$20"
ws_scenarios['D10'] = "=-'What''s Left at Month End'!$B$20"

ws_plan = wb['Monthly Plan']
ws_plan['B23'] = "=-'What''s Left at Month End'!$B$20"
ws_plan['B27'] = "=IF(F8=0,0,(B17+'What''s Left at Month End'!$B$20)/F8)"

wb.save(wb_path)
print("Workbook updated and saved successfully!")
