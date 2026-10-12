import openpyxl

wb = openpyxl.load_workbook('d:/cash_machine/Chathura_Printers_Business_Plan.xlsx')
print('Sheets in workbook:', wb.sheetnames)
ws = wb["What's Left at Month End"]
for r in [4, 5, 6, 8, 10, 11, 12, 13, 14, 18, 19, 20, 23, 24, 29, 34, 36, 38, 40, 41, 42, 44, 45]:
    vals = [f"{ws.cell(r, c).coordinate}={ws.cell(r, c).value}" for c in range(1, 7) if ws.cell(r, c).value is not None]
    print(f"Row {r:2d}: " + " | ".join(vals))
